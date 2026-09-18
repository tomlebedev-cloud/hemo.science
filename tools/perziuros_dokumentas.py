"""Sudaro Word dokumentą autoriui peržiūrėti konspektus ir gidus.

Į dokumentą patenka: bendri neatitikimai tarp pranešimų, kiekvieno pranešimo neaiškios
vietos (iš konspekto skilties „Ką dar pasitikrinti“, be vidinių pastabų apie failus)
ir rodmenų lentelės su tuščiu stulpeliu pataisoms.

Paleidimas: python tools/perziuros_dokumentas.py
Rezultatas: perziura/Konspektu perziura YYYY-MM-DD.docx (į repozitoriją nekeliama)
"""
import datetime
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / 'konspektavimas'
OUT_DIR = ROOT / 'perziura'
SITE = 'https://tomlebedev-cloud.github.io/hemo.science/'

# vidinės pastabos, kurios autoriui neaktualios
SKIP = re.compile(r'Vaizdo įrašas|YouTube|skaidres/|kadr|sąraše|Kitų pranešėjų|Data sename|įkeltas du kartus|'
                  r'terminai (vietomis )?iškraipyti|pasikeitė|Skaidrės `|aplanke|tai atitinka', re.I)

COMMON = [
    ('D-dimerų amžiaus pataisa',
     'Pranešimuose amžiaus pataisa nuskambėjo skirtingai (Nr. 7 — 650–1500 ng/ml vyresniems; Nr. 35 — skaičiai neaiškūs). '
     'Ar taikoma taisyklė „amžius × 10 nuo 50 metų“ (µg/l FEU)?'),
    ('Klopidogrelio PRU terapinis intervalas',
     'Skirtinguose pranešimuose: 85–282 (Nr. 8), 86–230 (Nr. 10), 85–230 (Nr. 23), 90–240 (Nr. 3), 98–230 (Nr. 33). '
     'Kurį intervalą nurodyti gide?'),
    ('Fibrino monomerų vienetai',
     'Įrašuose skamba µg/ml, mg/l ir µg/l. Giduose rašoma µg/ml (= mg/l): norma ~4–6, nėštumo profilaktikos riba ~25, trombozė ~30–40. Ar teisinga?'),
    ('MMMH anti-Xa riba prieš operaciją',
     'Nr. 3 ir Nr. 25 — mažiau 0,2 TV/ml; Nr. 28 — mažiau 0,4. Kuri riba teisinga ar nuo ko priklauso?'),
    ('TGAK koncentracija prieš operaciją',
     'Nurodoma mažiau 50 ng/ml, saugiau mažiau 30, neuroaksinei nejautrai — arti 0. Ar taip nurodyti gide?'),
    ('Aspirino ARU apatinė riba',
     'Nr. 23: 450–550 — optimalu, mažiau 450 — kraujavimo rizika. Kituose pranešimuose minima tik riba 550. Ar 450 riba taikytina?'),
]


def plain(html):
    return re.sub(r'<[^>]+>', '', html).replace('&lt;', '<').replace('&gt;', '>').replace('&amp;', '&')


def clean_md(text):
    text = re.sub(r'`([^`]*)`', r'\1', text)
    return text.replace('**', '').strip()


def review_items(md):
    m = re.search(r'^## Ką dar pasitikrinti\n(.*?)(?=^## |\Z)', md, re.S | re.M)
    if not m:
        return []
    items = [clean_md(l[2:]) for l in m.group(1).splitlines() if l.startswith('- ')]
    return [i for i in items if not SKIP.search(i)]


def add_link(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
                          is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '7A1F2B')
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rpr.append(color)
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement('w:t')
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def table(doc, head, rows, widths):
    t = doc.add_table(rows=1, cols=len(head))
    t.style = 'Table Grid'
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]
        c.text = ''
        r = c.paragraphs[0].add_run(h)
        r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
    return t


def main():
    items = {it['id']: it for it in json.loads((ROOT / 'data' / 'pranesimai.json').read_text(encoding='utf-8'))}
    guides = json.loads((ROOT / 'data' / 'gidai.json').read_text(encoding='utf-8'))

    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)

    doc.add_heading('hemo.science — konspektų ir gidų peržiūra', 0)
    today = datetime.date.today().isoformat()
    doc.add_paragraph(f'Parengta {today}.')
    doc.add_paragraph(
        'Svetainėje hemo.science paskelbti pranešimų konspektai ir teminiai gidai parengti pagal pranešimų įrašus. '
        'Kol neperžiūrėti, prie jų rodoma pastaba „autoriaus neperžiūrėta“.')
    doc.add_paragraph(
        'Neaiškios vietos — tai įrašo fragmentai, kuriuos automatinis kalbos atpažinimas užrašė netiksliai. Prašau peržiūrėti toliau pažymėtas vietas. Lentelėse stulpelyje „Pataisa“ užtenka įrašyti teisingą reikšmę; '
        'jei viskas teisinga — palikti tuščią. Kitas pastabas galima rašyti tiesiog tekste arba Word komentaruose.')
    p = doc.add_paragraph('Svetainė: ')
    add_link(p, SITE, SITE)

    doc.add_heading('1. Bendri klausimai', 1)
    doc.add_paragraph('Vietos, kur skirtinguose pranešimuose skaičiai nesutampa. Nuo atsakymų priklauso gidų turinys.')
    table(doc, ['Tema', 'Klausimas', 'Atsakymas'], COMMON, [4, 9.5, 4])

    doc.add_heading('2. Pranešimai', 1)
    n_items = n_rows = 0
    for folder in sorted(SRC.iterdir(), key=lambda f: f.name):
        k = folder / 'konspektas.md'
        if not (folder.is_dir() and k.exists()):
            continue
        nr = int(folder.name[:2])
        it = items[nr]
        data = json.loads((ROOT / 'data' / 'konspektai' / f'{nr}.json').read_text(encoding='utf-8'))
        questions = review_items(k.read_text(encoding='utf-8'))
        rows = [[plain(c) for c in r] for r in (data['table']['rows'] if data['table'] else [])]

        date = it['date'] or (str(it['year']) if it['year'] else '')
        doc.add_heading(f'Nr. {nr}. {it["title"]}', 2)
        p = doc.add_paragraph(f'{date} · ')
        add_link(p, f'{SITE}pranesimas.html?nr={nr}', 'konspektas svetainėje')
        if questions:
            doc.add_paragraph('Neaiškios vietos:').runs[0].bold = True
            for q in questions:
                doc.add_paragraph(q, style='List Bullet')
            n_items += len(questions)
        if rows:
            doc.add_paragraph('Rodmenys ir ribos konspekte:').runs[0].bold = True
            if len(rows[0]) == 2:
                table(doc, ['Rodmuo', 'Reikšmė konspekte', 'Pataisa'], [r + [''] for r in rows], [5.5, 7.5, 4.5])
            else:
                table(doc, ['Rodmuo', 'Tyrimas', 'Reikšmė konspekte', 'Pataisa'], [r[:3] + [''] for r in rows],
                      [4.5, 3, 6, 4])
            n_rows += len(rows)
        doc.add_paragraph()

    doc.add_heading('3. Gidai', 1)
    doc.add_paragraph(
        'Gidai sujungia kelių pranešimų medžiagą pagal temą. Kiekvienas teiginys turi nuorodą į pranešimą ir įrašo vietą. '
        'Juos verta perskaityti svetainėje ir pažymėti, ką taisyti.')
    for g in guides:
        p = doc.add_paragraph(style='List Bullet')
        add_link(p, f'{SITE}gidas.html?id={g["id"]}', g['title'])
        p.add_run(f' — {g["lead"]}')

    OUT_DIR.mkdir(exist_ok=True)
    out = OUT_DIR / f'Konspektu perziura {today}.docx'
    doc.save(out)
    print(f'{out.relative_to(ROOT)}: klausimų {n_items + len(COMMON)}, lentelių eilučių {n_rows}')


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
