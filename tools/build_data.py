"""Sugeneruoja data/pranesimai.json iš pranešimų sąrašo Excel failo (lapas „Puslapiui“).

Paleidimas:  python tools/build_data.py [kelias/iki/Pranešimų sąrašas.xlsx]
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = ROOT.parent / 'DS' / 'PRESENTATIONS' / 'Pranešimų sąrašas 2026-09-17.xlsx'
OUT = ROOT / 'data' / 'pranesimai.json'

LANG_CODE = {'Lietuvių': 'lt', 'Anglų': 'en', 'Rusų': 'ru', 'Ukrainiečių': 'uk'}
CONFERENCES = {'Konferencijos ir seminarai'}


def yt_id(url):
    m = re.search(r'(?:youtu\.be/|v=)([\w-]{11})', url or '')
    return m.group(1) if m else None


def seconds(td):
    return int(td.total_seconds()) if td else None


def main():
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    ws = openpyxl.load_workbook(xlsx, data_only=True)['Puslapiui']
    head = [c.value for c in ws[1]]
    col = {name: i for i, name in enumerate(head)}

    items = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not isinstance(row[col['Nr.']], int):
            continue
        title = row[col['Tema']].strip()
        date = row[col['Data']]
        date_src = row[col['Datos šaltinis']] or ''
        read_lang = row[col['Skaitymo kalba']]
        slides = row[col['Skaidrių kalba']]
        category = row[col['Kategorija']]
        items.append({
            'id': row[col['Nr.']],
            'title': title,
            'year': row[col['Metai']],
            # tikslią dieną rodome tik jei ji ne iš YouTube įkėlimo
            'date': date.strftime('%Y-%m-%d') if date and date_src != 'YouTube įkėlimas' else None,
            'category': category,
            'lang': LANG_CODE.get(read_lang),
            'slides': None if slides in (None, 'Nežinoma') else slides,
            'duration': seconds(row[col['Trukmė']]),
            'youtube': yt_id(row[col['YouTube nuoroda']]),
            'conference': category in CONFERENCES and 'Trombozės dienos' in title,
        })

    items.sort(key=lambda x: (x['year'] or 0, x['date'] or ''), reverse=True)
    OUT.write_text(json.dumps(items, ensure_ascii=False, indent=1), encoding='utf-8')
    with_video = sum(1 for x in items if x['youtube'])
    print(f'{len(items)} pranešimų ({with_video} su YouTube nuoroda) → {OUT.relative_to(ROOT)}')


if __name__ == '__main__':
    main()
